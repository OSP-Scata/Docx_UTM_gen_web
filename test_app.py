import os
import io
import pytest
from docx import Document
from fastapi.testclient import TestClient
from app import app, modify_url
import openpyxl
from zipfile import ZipFile

client = TestClient(app)

# =====================================================================
# 1. МОДУЛЬНЫЕ ТЕСТЫ ЛОГИКИ ССЫЛОК
# =====================================================================


@pytest.mark.parametrize('source_url, campaign, source, medium, content, term, expected', [
    ('https://site.ru', 'sale', 'vk', 'article', '', '',
     'https://site.ru?utm_source=vk&utm_medium=article&utm_campaign=sale'),
    ('https://site.ru?id=123', 'sale', 'tg', 'post', 'btn', '',
     'https://site.ru?id=123&utm_source=tg&utm_medium=post&utm_campaign=sale&utm_content=btn'),
])
def test_modify_url_logic(source_url, campaign, source, medium, content, term, expected):
    result = modify_url(source_url, campaign, source, medium, content, term)
    assert result == expected
    assert result.count('?') == 1


# =====================================================================
# 2. ИНТЕГРАЦИОННЫЙ ТЕСТ ПОЛУЧЕНИЯ ZIP-АРХИВА
# =====================================================================
def test_generate_utm_returns_zip():
    """Проверяем, что бэкенд принимает форму и возвращает ZIP-архив."""

    test_docx_path = 'temp_test_document.docx'
    doc = Document()
    doc.add_paragraph("Ссылка для теста: https://mysite.ru")
    doc.save(test_docx_path)

    form_data = {
        'campaign_name': 'test_campaign',
        'utm_medium': 'article',
        'custom_medium': '',
        'utm_content': '',
        'utm_term': '',
        'custom_platform': 'vk, telegram'
    }

    try:
        with open(test_docx_path, 'rb') as f:
            files = {
                'file': ('test.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            }
            response = client.post(
                '/generate-utm/', data=form_data, files=files)

        assert response.status_code == 200
        assert response.headers['content-type'] == 'application/zip'
        assert response.content.startswith(b'PK')

    finally:
        if os.path.exists(test_docx_path):
            os.remove(test_docx_path)

# =====================================================================
# 3. ТЕСТ МУЛЬТИЗАГРУЗКИ С EXCEL-МАНИФЕСТОМ
# =====================================================================


def test_generate_utm_batch_returns_zip():
    """Проверяем, что эндпоинт массовой загрузки принимает Excel + ZIP и возвращает готовый архив."""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Имя файла статьи', 'Кампания (utm_campaign)',
              'Источники', 'Тип рекламы', 'Содержимое', 'Ключевое слово'])
    ws.append(['article1.docx', 'test_batch_campaign',
              'vk, telegram', 'cpc', 'banner', 'smm'])
    excel_stream = io.BytesIO()
    wb.save(excel_stream)
    excel_stream.seek(0)
    doc = Document()
    doc.add_paragraph("Ссылка для проверки: https://mysite.ru")
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_stream.seek(0)

    zip_stream = io.BytesIO()
    with ZipFile(zip_stream, 'w') as zf:
        zf.writestr('article1.docx', docx_stream.read())
    zip_stream.seek(0)

    batch_files = {
        'excel_file': ('manifest.xlsx', excel_stream, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        'zip_file': ('articles.zip', zip_stream, 'application/zip')
    }

    response = client.post('/generate-utm-batch/', files=batch_files)

    assert response.status_code == 200
    assert response.headers['content-type'] == 'application/zip'
    assert response.content.startswith(b'PK')
